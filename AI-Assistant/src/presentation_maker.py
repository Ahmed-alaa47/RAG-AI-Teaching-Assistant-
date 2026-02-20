import os
import logging
from typing import Optional, List
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

logger = logging.getLogger(__name__)

class PresentationMaker:
    """Handles the creation of PowerPoint presentations with professional styling and user-provided images."""

    def __init__(self, output_dir="presentations"):
        self.output_dir = output_dir
        if not os.path.exists(self.output_dir):
            os.makedirs(self.output_dir)
        
        # Define brand colors
        self.COLOR_BG = RGBColor(0, 33, 71)      # Midnight Blue
        self.COLOR_TITLE = RGBColor(255, 255, 255) # White
        self.COLOR_BODY = RGBColor(220, 220, 220)  # Light Gray
        self.COLOR_ACCENT = RGBColor(255, 193, 7) # Yellow-Amber

    def create_presentation(self, slides_data, image_paths: List[str] = None, filename="presentation.pptx"):
        """
        Creates a professionally styled .pptx file using user-provided local images.
        """
        try:
            prs = Presentation()
            # Set slide size for widescreen (16:9)
            prs.slide_width = Inches(13.333)
            prs.slide_height = Inches(7.5)
            
            # Use provided image paths
            image_paths = image_paths or []
            
            for idx, slide_info in enumerate(slides_data):
                # Use Title and Content layout
                slide_layout = prs.slide_layouts[1]
                slide = prs.slides.add_slide(slide_layout)
                
                # 1. Set Background Style
                background = slide.background
                fill = background.fill
                fill.solid()
                fill.fore_color.rgb = self.COLOR_BG
                
                # 2. Add Top Accent Bar
                accent_bar = slide.shapes.add_shape(
                    MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.1)
                )
                accent_bar.fill.solid()
                accent_bar.fill.fore_color.rgb = self.COLOR_ACCENT
                accent_bar.line.fill.background() # No outline
                
                # 3. Determine if we have an image for this slide
                img_path = None
                if idx < len(image_paths):
                    img_path = image_paths[idx]

                # 4. Style Title
                title_shape = slide.shapes.title
                if title_shape:
                    title_shape.text = slide_info.get("title", "Untitled Slide")
                    # Adjust title position/width if we have an image
                    if img_path:
                        title_shape.left = Inches(0.5)
                        title_shape.width = Inches(12)
                    
                    title_tf = title_shape.text_frame
                    for paragraph in title_tf.paragraphs:
                        for run in paragraph.runs:
                            run.font.color.rgb = self.COLOR_TITLE
                            run.font.name = 'Arial'
                            run.font.size = Pt(44)
                            run.font.bold = True
                
                # 5. Style Content (Body)
                body_shape = slide.placeholders[1]
                if body_shape:
                    # Adjust position for two-column if image exists
                    if img_path:
                        body_shape.left = Inches(0.5)
                        body_shape.top = Inches(1.5)
                        body_shape.width = Inches(6.5) # Left side
                        body_shape.height = Inches(5.5)
                    else:
                        body_shape.left = Inches(1.0)
                        body_shape.top = Inches(1.8)
                        body_shape.width = Inches(11.33)
                        body_shape.height = Inches(5.0)

                    tf = body_shape.text_frame
                    tf.word_wrap = True
                    
                    content = slide_info.get("content", [])
                    if isinstance(content, str):
                        content = [content]
                    
                    tf.text = ""
                    for i, point in enumerate(content):
                        p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
                        p.text = point
                        p.level = 0
                        p.font.color.rgb = self.COLOR_BODY
                        p.font.name = 'Arial'
                        p.font.size = Pt(28 if img_path else 32) # Slightly smaller if image
                
                # 6. Insert Local Image if available
                if img_path and os.path.exists(img_path):
                    try:
                        # Right side image
                        left = Inches(7.5)
                        top = Inches(1.8)
                        width = Inches(5.0)
                        slide.shapes.add_picture(img_path, left, top, width=width)
                    except Exception as e:
                        logger.error(f"Failed to insert local image into slide {idx}: {e}")
            
            output_path = os.path.join(self.output_dir, filename)
            base, ext = os.path.splitext(output_path)
            counter = 1
            while os.path.exists(output_path):
                output_path = f"{base}_{counter}{ext}"
                counter += 1
                
            prs.save(output_path)
            logger.info(f"Presentation with local images saved to: {output_path}")
            return output_path
        except Exception as e:
            logger.error(f"Error creating presentation: {e}")
            import traceback
            logger.error(traceback.format_exc())
            return None
